"""
Generate a Word (.docx) copy of README.md → Tinder_AI_Coach_Setup_Guide.docx.

    pip install python-docx
    python tools/make_docx.py

Handles headings, paragraphs, bullet/numbered lists, fenced code blocks, tables, and blockquotes.
(python-docx is only needed to regenerate the doc; it's not a runtime dependency of the app.)
"""
import re
from pathlib import Path

from docx import Document
from docx.shared import Pt

ROOT = Path(__file__).resolve().parent.parent
README = ROOT / "README.md"
OUT = ROOT / "Tinder_AI_Coach_Setup_Guide.docx"


def strip_inline(s: str) -> str:
    s = re.sub(r"\[([^\]]+)\]\(#[^)]+\)", r"\1", s)          # anchor links -> text only
    s = re.sub(r"\[([^\]]+)\]\(([^)]+)\)", r"\1 (\2)", s)    # external links -> text (url)
    s = s.replace("**", "").replace("`", "")
    return s


def main() -> None:
    lines = README.read_text(encoding="utf-8").splitlines()
    doc = Document()
    doc.styles["Normal"].font.name = "Calibri"
    doc.styles["Normal"].font.size = Pt(11)

    i, n = 0, len(lines)
    while i < n:
        line = lines[i]

        if line.strip().startswith("```"):                  # code block
            i += 1
            code = []
            while i < n and not lines[i].strip().startswith("```"):
                code.append(lines[i]); i += 1
            i += 1
            run = doc.add_paragraph().add_run("\n".join(code))
            run.font.name = "Consolas"; run.font.size = Pt(9.5)
            continue

        if line.strip().startswith("|") and i + 1 < n and set(lines[i + 1].strip()) <= set("|-: "):
            rows = []
            while i < n and lines[i].strip().startswith("|"):
                rows.append([c.strip() for c in lines[i].strip().strip("|").split("|")]); i += 1
            header, body = rows[0], rows[2:]
            table = doc.add_table(rows=1, cols=len(header))
            try:
                table.style = "Light Grid Accent 1"
            except Exception:
                pass
            for j, h in enumerate(header):
                table.rows[0].cells[j].text = strip_inline(h)
            for r in body:
                cells = table.add_row().cells
                for j, c in enumerate(r):
                    if j < len(cells):
                        cells[j].text = strip_inline(c)
            doc.add_paragraph()
            continue

        m = re.match(r"^(#{1,6})\s+(.*)", line)
        if m:
            doc.add_heading(strip_inline(m.group(2)), level=min(len(m.group(1)), 4))
            i += 1; continue

        if line.strip() in ("---", "***"):
            i += 1; continue

        if line.strip().startswith(">"):
            text = strip_inline(line.strip().lstrip(">").strip())
            if text:
                doc.add_paragraph().add_run(text).italic = True
            i += 1; continue

        if re.match(r"^\s*[-*]\s+", line):
            doc.add_paragraph(strip_inline(re.sub(r"^\s*[-*]\s+", "", line)), style="List Bullet")
            i += 1; continue

        if re.match(r"^\s*\d+\.\s+", line):
            doc.add_paragraph(strip_inline(re.sub(r"^\s*\d+\.\s+", "", line)), style="List Number")
            i += 1; continue

        if not line.strip():
            i += 1; continue

        doc.add_paragraph(strip_inline(line))
        i += 1

    doc.save(OUT)
    print("wrote", OUT)


if __name__ == "__main__":
    main()
